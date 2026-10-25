import unittest
from pathlib import Path
import tempfile
import io
import pypdf
import docx
import pptx

from src.document_loader import DocumentLoader
from src.text_splitter import DocumentSplitter
from src.vector_store import VectorStoreManager
from src.rag_pipeline import RAGPipeline
from langchain_core.documents import Document

class TestRAGPipeline(unittest.TestCase):
    
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_01_pdf_extraction(self):
        # Generate a test PDF with 2 pages
        pdf_writer = pypdf.PdfWriter()
        pdf_writer.add_blank_page(width=100, height=100)
        pdf_writer.add_blank_page(width=100, height=100)
        
        pdf_buffer = io.BytesIO()
        pdf_writer.write(pdf_buffer)
        pdf_buffer.seek(0)
        
        # Test loader (blank pages should yield empty or valid metadata handling)
        docs = DocumentLoader.load_pdf(pdf_buffer, "sample.pdf")
        self.assertIsInstance(docs, list)

    def test_02_docx_extraction(self):
        # Generate a test DOCX
        doc = docx.Document()
        doc.add_heading("Project Overview", level=1)
        doc.add_paragraph("This is the first section of the document about RAG pipelines.")
        doc.add_page_break()
        doc.add_heading("Technical Architecture", level=1)
        doc.add_paragraph("The technical architecture uses LangChain and FAISS for vector storage.")
        
        docx_buffer = io.BytesIO()
        doc.save(docx_buffer)
        docx_buffer.seek(0)
        
        docs = DocumentLoader.load_docx(docx_buffer, "specs.docx")
        self.assertGreaterEqual(len(docs), 1)
        self.assertEqual(docs[0].metadata["source"], "specs.docx")
        self.assertEqual(docs[0].metadata["doc_type"], "docx")
        self.assertIn("RAG pipelines", docs[0].page_content)

    def test_03_pptx_extraction(self):
        # Generate a test PPTX with 2 slides
        prs = pptx.Presentation()
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "Introduction to Retrieval"
        slide1.placeholders[1].text = "Retrieval Augmented Generation combines search and generation."

        slide2 = prs.slides.add_slide(prs.slide_layouts[0])
        slide2.shapes.title.text = "FAISS Vector Index"
        slide2.placeholders[1].text = "FAISS provides fast similarity search on dense embeddings."

        pptx_buffer = io.BytesIO()
        prs.save(pptx_buffer)
        pptx_buffer.seek(0)

        docs = DocumentLoader.load_pptx(pptx_buffer, "presentation.pptx")
        self.assertEqual(len(docs), 2)
        self.assertEqual(docs[0].metadata["page"], 1)
        self.assertEqual(docs[1].metadata["page"], 2)
        self.assertEqual(docs[0].metadata["source"], "presentation.pptx")
        self.assertIn("Retrieval Augmented Generation", docs[0].page_content)
        self.assertIn("FAISS", docs[1].page_content)

    def test_04_text_chunking_and_metadata_preservation(self):
        raw_docs = [
            Document(
                page_content="Artificial Intelligence and Machine Learning are transforming modern search engines. " * 30,
                metadata={"source": "ai_report.pdf", "page": 3, "doc_type": "pdf"}
            )
        ]
        splitter = DocumentSplitter(chunk_size=200, chunk_overlap=40)
        chunks = splitter.split_documents(raw_docs)
        self.assertGreater(len(chunks), 1)
        for chunk in chunks:
            self.assertEqual(chunk.metadata["source"], "ai_report.pdf")
            self.assertEqual(chunk.metadata["page"], 3)
            self.assertIn("chunk_id", chunk.metadata)

    def test_05_faiss_vector_store_and_persistence(self):
        test_chunks = [
            Document(
                page_content="FAISS is a library for efficient similarity search and clustering of dense vectors.",
                metadata={"source": "faiss_docs.pdf", "page": 1, "doc_type": "pdf", "chunk_id": "c1"}
            ),
            Document(
                page_content="Sentence-Transformers provides state of the art models for sentence and document embeddings.",
                metadata={"source": "st_docs.docx", "page": 2, "doc_type": "docx", "chunk_id": "c2"}
            ),
            Document(
                page_content="Streamlit turns data scripts into shareable web apps in minutes with pure Python.",
                metadata={"source": "streamlit.pptx", "page": 1, "doc_type": "pptx", "chunk_id": "c3"}
            ),
        ]

        vector_manager = VectorStoreManager(storage_dir=self.temp_path)
        vector_manager.create_vector_store(test_chunks)
        
        # Test similarity search
        results = vector_manager.similarity_search("How do dense embeddings get generated?", k=1)
        self.assertEqual(len(results), 1)
        self.assertIn("Sentence-Transformers", results[0].page_content)
        self.assertEqual(results[0].metadata["source"], "st_docs.docx")

        # Test index persistence
        vector_manager.save_index(self.temp_path)
        new_manager = VectorStoreManager(storage_dir=self.temp_path)
        loaded = new_manager.load_index(self.temp_path)
        self.assertTrue(loaded)

        loaded_results = new_manager.similarity_search("web apps in python", k=1)
        self.assertEqual(len(loaded_results), 1)
        self.assertIn("Streamlit", loaded_results[0].page_content)

    def test_06_rag_pipeline_answer_and_citations(self):
        test_chunks = [
            Document(
                page_content="Antigravity 2.0 was announced with major autonomous capabilities.",
                metadata={"source": "announcement.pdf", "page": 4, "doc_type": "pdf", "chunk_id": "a4"}
            )
        ]
        vector_manager = VectorStoreManager(storage_dir=self.temp_path)
        vector_manager.create_vector_store(test_chunks)

        # Initialize pipeline in hub mode without token to test graceful fallback & citation extraction
        pipeline = RAGPipeline(vector_store_manager=vector_manager, llm_type="huggingface_hub", hf_api_token="")
        response = pipeline.answer_query("When was Antigravity announced?")

        self.assertIn("sources", response)
        self.assertEqual(len(response["sources"]), 1)
        self.assertEqual(response["sources"][0]["source"], "announcement.pdf")
        self.assertEqual(response["sources"][0]["page"], 4)
        self.assertIn("Antigravity 2.0", response["sources"][0]["content"])

if __name__ == "__main__":
    unittest.main()
